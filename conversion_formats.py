"""Helpers for standard Markdown, PDF and Word exports."""

from __future__ import annotations

import re
from pathlib import Path
from urllib.parse import unquote, urlparse

import fitz
import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

MARKDOWN_EXTENSIONS = ["extra", "sane_lists", "smarty"]

PDF_CSS = """
@page { size: a4; margin: 0; }
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.48; color: #172033; }
h1, h2, h3, h4, h5, h6 { color: #172033; font-weight: bold; margin-top: 1.15em; margin-bottom: 0.45em; }
h1 { font-size: 24pt; border-bottom: 1px solid #ccd5e3; padding-bottom: 7px; }
h2 { font-size: 18pt; border-bottom: 1px solid #e3e8ef; padding-bottom: 4px; }
h3 { font-size: 14pt; }
p { margin: 0.45em 0 0.7em; }
a { color: #2457c5; text-decoration: none; }
blockquote { margin: 0.8em 0; padding: 8px 12px; border-left: 4px solid #7c5cff; background: #f5f2ff; color: #30394d; }
pre { background: #101827; color: #edf2f7; padding: 10px; border-radius: 5px; font-family: monospace; font-size: 8.5pt; white-space: pre-wrap; }
code { background: #eef1f5; color: #9b2743; padding: 1px 3px; font-family: monospace; }
pre code { background: transparent; color: inherit; padding: 0; }
table { border-collapse: collapse; width: 100%; margin: 0.8em 0; font-size: 9pt; }
th { background: #e9edf5; font-weight: bold; }
th, td { border: 1px solid #b9c2d0; padding: 5px 7px; vertical-align: top; }
img { max-width: 100%; height: auto; }
hr { border: 0; border-top: 1px solid #ccd5e3; margin: 1.2em 0; }
"""


def markdown_to_html(markdown_text: str) -> str:
    """Render portable standard Markdown as HTML."""
    standard_text = markdown_text.replace("\r\n", "\n").strip() + "\n"
    return markdown.markdown(
        standard_text,
        extensions=MARKDOWN_EXTENSIONS,
        output_format="html5",
    )

def _safe_image_path(src: str, base_dir: Path) -> Path | None:
    parsed = urlparse(src)
    if parsed.scheme or not src:
        return None
    decoded = unquote(parsed.path).replace("\\", "/")
    base = base_dir.resolve()
    for candidate in (base_dir / decoded, base_dir / Path(decoded).name):
        try:
            resolved = candidate.resolve()
            resolved.relative_to(base)
        except (ValueError, OSError):
            continue
        if resolved.is_file():
            return resolved
    return None


def _prepare_html_assets(html_text: str, base_dir: Path) -> str:
    soup = BeautifulSoup(html_text, "html.parser")
    for image in soup.find_all("img"):
        src = image.get("src", "")
        local = _safe_image_path(src, base_dir)
        if local is not None:
            image["src"] = local.name
        elif not re.match(r"^https?://", src, re.I):
            replacement = soup.new_tag("span")
            replacement.string = f"[Imagen no disponible: {image.get('alt') or src}]"
            image.replace_with(replacement)
    return str(soup)


def markdown_to_pdf(markdown_text: str, output_path: Path, base_dir: Path, title: str) -> Path:
    """Render standard Markdown to a paginated searchable PDF."""
    html_body = _prepare_html_assets(markdown_to_html(markdown_text), base_dir)
    html_doc = "<html><head><meta charset='utf-8'></head><body><article>" + html_body + "</article></body></html>"
    story = fitz.Story(html=html_doc, user_css=PDF_CSS, archive=fitz.Archive(str(base_dir)))
    page_rect = fitz.paper_rect("a4")
    content_rect = page_rect + (50, 54, -50, -54)

    def rect_function(_rect_num, _filled):
        return page_rect, content_rect, None

    pdf = story.write_with_links(rect_function)
    metadata = pdf.metadata or {}
    metadata.update({"title": title, "producer": "PDF2MD Batch Converter"})
    pdf.set_metadata(metadata)
    pdf.save(str(output_path))
    pdf.close()
    return output_path


def _configure_document(document: Document, title: str):
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    for level in range(1, 7):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.color.rgb = RGBColor(23, 32, 51)
    document.core_properties.title = title
    document.core_properties.author = "PDF2MD Batch Converter"


def _add_inline(paragraph, node, bold=False, italic=False, code=False):
    if isinstance(node, NavigableString):
        if str(node):
            run = paragraph.add_run(str(node))
            run.bold, run.italic = bold, italic
            if code:
                run.font.name, run.font.size = "Consolas", Pt(9)
                run.font.color.rgb = RGBColor(155, 39, 67)
        return
    if not isinstance(node, Tag):
        return
    name = node.name.lower()
    if name == "br":
        paragraph.add_run().add_break()
        return
    if name == "img":
        paragraph.add_run(f"[{node.get('alt') or node.get('src') or 'Imagen'}]")
        return
    if name == "a":
        label, href = node.get_text(" ", strip=True) or node.get("href", ""), node.get("href", "")
        run = paragraph.add_run(label)
        run.font.color.rgb, run.underline = RGBColor(36, 87, 197), True
        if href and href != label:
            tail = paragraph.add_run(f" ({href})")
            tail.font.size, tail.font.color.rgb = Pt(8), RGBColor(90, 99, 116)
        return
    for child in node.children:
        _add_inline(
            paragraph,
            child,
            bold or name in {"strong", "b"},
            italic or name in {"em", "i"},
            code or name == "code",
        )


def _add_image(document: Document, tag: Tag, base_dir: Path) -> bool:
    image_path = _safe_image_path(tag.get("src", ""), base_dir)
    if image_path is None:
        return False
    try:
        document.add_picture(str(image_path), width=Inches(6.35))
        document.paragraphs[-1].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        if tag.get("alt"):
            caption = document.add_paragraph(tag["alt"])
            caption.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            for run in caption.runs:
                run.italic, run.font.size = True, Pt(8)
        return True
    except Exception:
        return False


def _render_list(document: Document, tag: Tag, ordered: bool, level: int = 0):
    for item in tag.find_all("li", recursive=False):
        paragraph = document.add_paragraph(style="List Number" if ordered else "List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.22 * level)
        for child in item.children:
            if not (isinstance(child, Tag) and child.name in {"ul", "ol"}):
                _add_inline(paragraph, child)
        for nested in item.find_all(["ul", "ol"], recursive=False):
            _render_list(document, nested, nested.name == "ol", level + 1)


def _render_table(document: Document, table_tag: Tag):
    rows = table_tag.find_all("tr")
    if not rows:
        return
    width = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_index, row_tag in enumerate(rows):
        for col_index, cell_tag in enumerate(row_tag.find_all(["th", "td"], recursive=False)):
            paragraph = table.cell(row_index, col_index).paragraphs[0]
            for child in cell_tag.children:
                _add_inline(paragraph, child, bold=cell_tag.name == "th")


def markdown_to_docx(markdown_text: str, output_path: Path, base_dir: Path, title: str) -> Path:
    """Render standard Markdown to a styled editable Word file."""
    soup = BeautifulSoup(markdown_to_html(markdown_text), "html.parser")
    document = Document()
    _configure_document(document, title)
    for node in soup.children:
        if isinstance(node, NavigableString):
            if str(node).strip():
                document.add_paragraph(str(node).strip())
            continue
        if not isinstance(node, Tag):
            continue
        name = node.name.lower()
        if re.fullmatch(r"h[1-6]", name):
            paragraph = document.add_heading(level=int(name[1]))
            for child in node.children:
                _add_inline(paragraph, child)
        elif name == "p":
            images = node.find_all("img")
            if node.get_text(" ", strip=True):
                paragraph = document.add_paragraph()
                for child in node.children:
                    if not (isinstance(child, Tag) and child.name == "img"):
                        _add_inline(paragraph, child)
            for image in images:
                if not _add_image(document, image, base_dir):
                    document.add_paragraph(f"[Imagen no disponible: {image.get('alt') or image.get('src', '')}]")
        elif name in {"ul", "ol"}:
            _render_list(document, node, name == "ol")
        elif name == "blockquote":
            paragraph = document.add_paragraph(style="Quote")
            for child in node.children:
                _add_inline(paragraph, child)
        elif name == "pre":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            run = paragraph.add_run(node.get_text())
            run.font.name, run.font.size = "Consolas", Pt(8.5)
            run.font.color.rgb = RGBColor(35, 48, 68)
        elif name == "table":
            _render_table(document, node)
        elif name == "hr":
            document.add_paragraph("────────────────────────────────────────")
        elif name == "img":
            _add_image(document, node, base_dir)
        else:
            paragraph = document.add_paragraph()
            for child in node.children:
                _add_inline(paragraph, child)
    document.save(str(output_path))
    return output_path
